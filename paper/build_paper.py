"""Build an editable IEEE-style manuscript and original scientific figures."""
import gzip
import json
import math
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, Rectangle
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
S = json.loads((ROOT/'results/summary.json').read_text())
R = {(r['scenario'], r['architecture']): r for r in S['results']}
A = ['N', 'N+1', '2N']
BLUE, TEAL, ORANGE = '#244a70', '#287e83', '#b45f2c'
COLORS = [BLUE, TEAL, ORANGE]
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':8,'axes.labelsize':8,
                     'xtick.labelsize':8,'ytick.labelsize':7,'legend.fontsize':7,
                     'axes.spines.top':False,'axes.spines.right':False,'svg.fonttype':'none'})


def save(fig, name):
    fig.savefig(ROOT/'figures'/f'{name}.png', dpi=450, bbox_inches='tight', pad_inches=.05)
    fig.savefig(ROOT/'figures'/f'{name}.svg', bbox_inches='tight', pad_inches=.05)
    plt.close(fig)


def box(ax, x, y, w, h, text, fc='#eef3f6', fs=8):
    ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.015,rounding_size=0.035',
                              edgecolor=BLUE,facecolor=fc,linewidth=.85))
    ax.text(x+w/2,y+h/2,text,ha='center',va='center',fontsize=fs,color='#172b3a')


def arrow(ax,x1,y1,x2,y2):
    ax.annotate('',xy=(x2,y2),xytext=(x1,y1),arrowprops={'arrowstyle':'->','color':BLUE,'lw':1})


def figures():
    fig,ax=plt.subplots(figsize=(3.45,2.75)); ax.set(xlim=(0,10),ylim=(0,10));ax.axis('off')
    box(ax,1,8.5,8,1.1,'Facility inputs: load, ratings, topology')
    box(ax,1,6.55,8,1.15,'Integer sizing + complete A/B trains')
    box(ax,1,4.55,8,1.2,'Seeded failure / repair event streams')
    box(ax,1,2.55,8,1.2,'Chronological capacity integration')
    box(ax,1,.5,8,1.2,'Availability, SLA risk, paired comparison')
    for top,bottom in [(8.5,7.7),(6.55,5.75),(4.55,3.75),(2.55,1.7)]:arrow(ax,5,top,5,bottom)
    ax.text(9.55,5.1,'Versioned rates + scenario assumptions',rotation=90,ha='center',va='center',fontsize=7)
    save(fig,'fig1_workflow')

    fig,ax=plt.subplots(figsize=(3.45,2.35));ax.set(xlim=(0,10),ylim=(0,7));ax.axis('off')
    box(ax,.3,5.35,9.4,1.15,'N: one train; each group has n units')
    box(ax,.3,3.65,9.4,1.15,'N+1: one train; each group has n + 1 units')
    box(ax,.3,1.5,3.8,1.2,'2N: train A\nn units / group',fs=7.7)
    box(ax,.3,.05,3.8,1.2,'2N: train B\nn units / group',fs=7.7)
    box(ax,6.15,.65,3.55,1.35,'Select complete\nusable train',fs=7.7)
    arrow(ax,4.1,2.1,6.15,1.65);arrow(ax,4.1,.65,6.15,1.05)
    ax.text(5,3.2,'No cross-ties or pooled A/B capacity',ha='center',fontsize=7.5,color=ORANGE)
    save(fig,'fig2_topology')

    fig,axs=plt.subplots(2,1,figsize=(3.45,2.65),sharex=True,gridspec_kw={'height_ratios':[1,1.25]})
    axs[0].broken_barh([(1,4)],(.7,.45),facecolors=BLUE)
    axs[0].broken_barh([(3,4)],(.0,.45),facecolors=ORANGE)
    axs[0].set(yticks=[.22,.92],yticklabels=['UPS A02','UPS A01'],ylim=(-.1,1.4),xlim=(0,8))
    axs[0].spines[['bottom','left']].set_visible(False);axs[0].tick_params(axis='y',length=0)
    axs[1].step([0,1,3,5,7,8],[10,10,7.5,10,10,10],where='post',color=BLUE,lw=1.6)
    axs[1].fill_between([3,5],[7.5,7.5],[10,10],step='post',color=ORANGE,alpha=.18)
    axs[1].axhline(10,color='#555555',ls='--',lw=.7)
    axs[1].text(4,8.5,'2 h outage',ha='center',fontsize=7)
    axs[1].set(ylabel='Served load (MW)',xlabel='Time (h)',ylim=(7,10.6),yticks=[7.5,10])
    fig.tight_layout(h_pad=.5);save(fig,'fig3_overlap')

    fig,ax=plt.subplots(figsize=(3.45,2.5));x=np.arange(3);w=.32
    for j,(scenario,label,color,hatch) in enumerate([('baseline','Nominal rates',BLUE,''),('stress20','20× rate stress',ORANGE,'//')]):
        rows=[R[(scenario,a)] for a in A];ys=[r['downtime_minutes'] for r in rows]
        ax.bar(x+(j-.5)*w,ys,w,label=label,color=color,hatch=hatch,edgecolor='white')
        ax.errorbar(x+(j-.5)*w,ys,yerr=[[r['downtime_minutes']-r['ci_low'] for r in rows],[r['ci_high']-r['downtime_minutes'] for r in rows]],fmt='none',color='black',capsize=2,lw=.8)
    ax.set(xticks=x,xticklabels=A,yscale='log',ylabel='Mean downtime (min/year)',ylim=(1,150000))
    ax.legend(loc='upper right',frameon=False);ax.grid(axis='y',alpha=.18);ax.set_axisbelow(True)
    fig.tight_layout();save(fig,'fig4_comparison')

    fig,ax=plt.subplots(figsize=(3.45,2.5));x=np.arange(3);w=.23
    for j,(sc,label,color,hatch) in enumerate([('baseline','Independent',BLUE,''),('shared_bank','Bank-A hazard',TEAL,'//'),('shared_site','Site-wide hazard',ORANGE,'xx')]):
        vals=[R[(sc,a)]['downtime_minutes'] for a in A]
        ax.bar(x+(j-1)*w,vals,w,label=label,color=color,hatch=hatch,edgecolor='white')
    ax.set(xticks=x,xticklabels=A,yscale='log',ylabel='Mean downtime (min/year)',ylim=(1,12000))
    ax.legend(frameon=False,fontsize=6.8,loc='upper right');ax.grid(axis='y',alpha=.18);ax.set_axisbelow(True)
    fig.tight_layout();save(fig,'fig5_dependence')

    fig,ax=plt.subplots(figsize=(3.45,2.35));x=np.arange(3);w=.32
    for j,(sc,label,color) in enumerate([('baseline','Printed rate: 0.58/year',BLUE),('generator_alternative','Count/exposure: 0.4323/year',TEAL)]):
        rows=[R[(sc,a)] for a in A];y=[r['downtime_minutes'] for r in rows]
        ax.bar(x+(j-.5)*w,y,w,label=label,color=color,edgecolor='white',hatch='//' if j else '')
        ax.errorbar(x+(j-.5)*w,y,yerr=[[r['downtime_minutes']-r['ci_low'] for r in rows],[r['ci_high']-r['downtime_minutes'] for r in rows]],fmt='none',color='black',capsize=2,lw=.8)
    ax.set(xticks=x,xticklabels=A,yscale='log',ylabel='Mean downtime (min/year)',ylim=(.6,18000))
    ax.legend(frameon=False,fontsize=6.6,loc='upper right');ax.grid(axis='y',alpha=.18);ax.set_axisbelow(True)
    fig.tight_layout();save(fig,'fig6_source_audit')

    fig,axs=plt.subplots(1,2,figsize=(3.45,2.2),sharey=True)
    for ax,arch in zip(axs,['N+1','2N']):
        for seed,color in zip(S['metadata']['seeds'],COLORS):
            with gzip.open(ROOT/'results'/f'baseline_{arch.replace("+","plus")}_{seed}.json.gz','rt') as f:r=json.load(f)['result']
            vals=np.array([t['annual_downtime_minutes'] for t in r['trial_results']]);ks=np.arange(1000,10001,250)
            ax.plot(ks/1000,[vals[:k].mean() for k in ks],lw=1,color=color,label=str(seed))
        ax.set(title=arch,xlabel='Trials (thousands)',xticks=[2,6,10],ylim=(0,12));ax.grid(alpha=.18)
    axs[0].set_ylabel('Cumulative downtime (min/year)')
    axs[1].legend(title='Seed',frameon=False,fontsize=6,title_fontsize=6,loc='upper right')
    fig.tight_layout(w_pad=.6);save(fig,'fig7_convergence')

D=Document()
sec=D.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.75);sec.bottom_margin=Inches(1)
sec.left_margin=sec.right_margin=Inches(.625)
sec.header_distance=sec.footer_distance=Inches(.3)
style=D.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(10)
style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
style.paragraph_format.first_line_indent=Inches(.14)
style.paragraph_format.space_after=Pt(0)
style.paragraph_format.line_spacing=1.0
style.paragraph_format.widow_control=True
for name in ['Title','Heading 1','Heading 2','Caption']:
    st=D.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
D.core_properties.title='DC-Resilience: A Monte Carlo Failure-Injection Framework for Quantitative Evaluation of Data Centre Redundancy Architectures'
D.core_properties.author='Romit Deokar; Shourya Saran'
D.core_properties.subject='Reproducible, assumption-based infrastructure availability study'
D.core_properties.keywords='Monte Carlo; data centre; redundancy; availability; failure injection'
settings=D.settings.element
compat=settings.find(qn('w:compat'))
if compat is not None:
    node=OxmlElement('w:compatSetting');node.set(qn('w:name'),'compatibilityMode');node.set(qn('w:uri'),'http://schemas.microsoft.com/office/word');node.set(qn('w:val'),'15');compat.append(node)


def p(text='',bold=False,italic=False,size=None,align=None,indent=True):
    z=D.add_paragraph();z.paragraph_format.first_line_indent=Inches(.14 if indent else 0)
    if align is not None:z.alignment=align
    r=z.add_run(text);r.bold=bold;r.italic=italic
    if size:r.font.size=Pt(size)
    return z


def h(text,sub=False):
    z=p(text,italic=sub,align=WD_ALIGN_PARAGRAPH.LEFT if sub else WD_ALIGN_PARAGRAPH.CENTER,indent=False)
    z.paragraph_format.space_before=Pt(7 if sub else 9)
    z.paragraph_format.space_after=Pt(3)
    z.paragraph_format.keep_with_next=True
    if not sub:
        z.runs[0].font.small_caps=True
    return z


def eq(text,num):
    z=p(indent=False);z.alignment=WD_ALIGN_PARAGRAPH.CENTER
    z.paragraph_format.space_before=Pt(5);z.paragraph_format.space_after=Pt(5)
    mathnode=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);mathnode.append(r);z._p.append(mathnode)
    z.add_run('    ('+str(num)+')').font.size=Pt(9)


def fig(name,caption):
    z=p(indent=False,align=WD_ALIGN_PARAGRAPH.CENTER)
    z.paragraph_format.space_before=Pt(5);z.paragraph_format.keep_with_next=True
    z.add_run().add_picture(str(ROOT/'figures'/f'{name}.png'),width=Inches(3.43))
    c=p(caption,size=8,indent=False);c.paragraph_format.space_after=Pt(6)
    c.paragraph_format.keep_together=True


def table(number,title,headers,rows,widths=None,note=None):
    cap=p('TABLE '+number+'\n'+title.upper(),size=8,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
    cap.paragraph_format.space_before=Pt(6);cap.paragraph_format.space_after=Pt(3);cap.paragraph_format.keep_with_next=True
    t=D.add_table(rows=1, cols=len(headers));t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trPr.append(no)
        if ri==0: rep=OxmlElement('w:tblHeader');trPr.append(rep)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            tcPr=c._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'42' if side in ['top','bottom'] else '45');el.set(qn('w:type'),'dxa');marg.append(el)
            tcPr.append(marg)
            for para in c.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.space_after=Pt(0)
                para.paragraph_format.keep_with_next=ri<len(t.rows)-1
                for run in para.runs:run.font.name='Times New Roman';run.font.size=Pt(8);run.bold=ri==0
            borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom']:
                if (side=='top' and ri==0) or (side=='bottom' and ri in [0,len(t.rows)-1]):
                    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');b.set(qn('w:color'),'333333');borders.append(b)
            tcPr.append(borders)
    if note:
        z=p(note,size=8,indent=False);z.paragraph_format.space_before=Pt(2);z.paragraph_format.space_after=Pt(5)
    else:
        z=p('',size=2);z.paragraph_format.space_after=Pt(1)


def val(sc,a,key='downtime_minutes',digits=2):return f'{R[(sc,a)][key]:,.{digits}f}'


def manuscript():
    title=p(D.core_properties.title,size=24,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
    title.paragraph_format.space_after=Pt(10);title.paragraph_format.line_spacing=1.0
    p('Romit Deokar and Shourya Saran',size=11,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
    p('Department of Computing Technologies, School of Computing',size=10,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
    p('SRM Institute of Science and Technology, Kattankulathur, Tamil Nadu, India',size=10,align=WD_ALIGN_PARAGRAPH.CENTER,indent=False)
    s=D.add_section(WD_SECTION_START.CONTINUOUS)
    cols=s._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'2');cols.set(qn('w:space'),'360')
    p('Abstract—Redundancy labels alone do not quantify the interruption risk of a data centre. This paper presents DC-Resilience, a continuous-time Monte Carlo framework that compares capacity-based N, N+1, and 2N architectures while retaining explicit failure-rate provenance, event traces, and statistical uncertainty. The model integrates overlapping failure and repair intervals and distinguishes pooled spare capacity from complete, disconnected power trains. An infrastructure-only case study holds a 10 MW load and equipment ratings constant across 450,000 simulated facility-years. Three independent seeds are used for nominal, accelerated-hazard, common-cause, and source-sensitivity experiments. Under nominal independent hazards, mean annual downtime is '+val('baseline','N')+', '+val('baseline','N+1')+', and '+val('baseline','2N')+' minutes for N, N+1, and 2N, respectively. The nominal difference between the two redundant designs is inconclusive, whereas bank-local and site-wide shared hazards produce distinct rankings. A generator-source inconsistency materially changes absolute predictions. The contribution is a reproducible framework and a transparent scenario study, not a field-calibrated forecast: two of the five active component classes use explicitly illustrative inputs, and ideal repair and failover assumptions remain. Results demonstrate why topology, dependency scope, input evidence, and confidence limits must accompany availability claims.',bold=True,size=9,indent=False)
    z=p('Index Terms—Data centre availability, failure injection, Monte Carlo simulation, redundancy, common-cause failures, reliability engineering.',bold=True,italic=True,size=9,indent=False)
    z.paragraph_format.space_before=Pt(5)

    h('I. Introduction')
    p('Data centres depend on coordinated power conversion, electrical distribution, and cooling to sustain an IT workload. Installing spare units can prevent an individual equipment failure from becoming a service interruption. Nevertheless, a capacity label such as N+1 does not specify every distribution path, repair policy, shared hazard, or operating condition that determines system availability. An engineering comparison therefore requires an explicit topology and a defensible stochastic model rather than a direct lookup from redundancy to uptime.')
    p('Uptime Institute classifies site infrastructure by performance characteristics including redundant capacity, concurrent maintainability, and fault tolerance [1]. It does not certify a fixed annual uptime percentage; references to expected annual downtime were removed from its Tier Standard in 2009 [2]. Accordingly, this work treats 99.982% only as a selected service-level agreement (SLA) benchmark. No simulated result is interpreted as Tier certification, and neither N+1 nor 2N is equated with a particular Tier.')
    p('The practical question is not simply whether more equipment improves resilience. It is whether a specified arrangement sustains the same demand under independent failures, overlapping repairs, and hazards that defeat multiple units together. This distinction matters when a single pooled spare can compensate for any failed unit within a group but a dual-train design cannot combine partial capacity across disconnected paths.')
    p('DC-Resilience addresses this question through a reproducible, event-driven comparison. Its contributions are: (1) a capacity-aware architecture model that prevents invalid A/B pooling; (2) exact integration of interruption duration and unserved load over failure/repair timelines; (3) paired, multi-seed experiments with explicit input provenance and uncertainty; and (4) a scenario study showing how failure dependence and source ambiguity affect the apparent value of redundancy. The framework implements established reliability methods rather than claiming a new Monte Carlo algorithm.')

    h('II. Related Work')
    p('IEEE Std 493-2007 provides reliability-engineering guidance and equipment statistics for industrial and commercial power systems [3]. Its use requires matching equipment class, operating duty, and statistical units to the proposed application. Wiboonrat [4] reproduces selected IEEE-493-attributed entries in a data-centre maintenance study. Those entries support an auditable secondary-source baseline, but do not independently establish failure rates for arbitrary modern modules, cooling equipment, or IT services.')
    p('Lei and Huang [5] developed a Monte Carlo next-event tool for data-centre power-distribution reliability. Their work establishes a direct methodological precedent. The present implementation emphasizes an inspectable web workflow, capacity-preserving comparisons, paired random streams, source-audit scenarios, and explicit separation between sampled outcomes and real-world forecasts. It does not claim precedence over next-event simulation or replace detailed electrical-protection models.')
    p('Ford et al. [6] studied availability in globally distributed storage systems and highlighted the importance of failure domains and correlated behaviour. Their storage-system observations motivate careful dependency modelling, but are not used as power-equipment rates. Julitz et al. [7] investigated dependent M-out-of-N architectures using Monte Carlo simulation and showed that the effect of dependence is architecture-specific. Here, dependence is deliberately simplified to a bank-local or site-wide shared power hazard; it is an interpretable stress mechanism, not a calibrated correlation model.')
    p('The statistical treatment also follows established practice. Exact binomial intervals are computed using Clopper–Pearson limits [8], and zero observed outages are not interpreted as proof that an event is impossible [9]. The gap addressed is therefore one of reproducible integration and transparent interpretation: topology, data provenance, scenario assumptions, and uncertainty are reported together for one controlled infrastructure case study.')

    h('III. System Model and Simulation Method')
    h('A. Capacity-Preserving Redundancy',True)
    p('Let L denote the required IT-support capacity in kW, and c_g the capacity of one unit in component group g. Groups comprise UPS modules, generators, power distribution units (PDUs), automatic transfer switches (ATSs), and computer-room air-conditioning (CRAC) units. Integer sizing is essential: the number of required units is')
    eq('n_g = ⌈L / c_g⌉',1)
    p('An N group contains n_g units, N+1 contains n_g+1 units in one train, and 2N contains two complete trains A and B, each with n_g units. An extra unit is continuously exposed to the same operating hazard as other units; standby failure rates and load-dependent hazards are not modelled. Fig. 1 summarizes the workflow, while Fig. 2 distinguishes the capacity arrangements.')
    fig('fig1_workflow','Fig. 1. DC-Resilience workflow. The same versioned inputs and scenario definitions drive the event engine and exported comparisons.')
    p('If H_g,b(t) units are healthy in group g and bank b, the bank capacity is c_g H_g,b(t). In the islanded operating mode, a usable power train requires UPS, generator, PDU, and ATS capacity simultaneously. Its capacity is the minimum across those groups. A 2N power subsystem selects the stronger complete train; it cannot combine UPS capacity from A with distribution capacity from B:')
    eq('C_power(t) = max_b min_g∈P {c_g H_g,b(t)}',2)
    p('Cooling capacity is selected from a complete cooling bank in the same way. The effective served capacity is')
    eq('C(t) = min {L, C_power(t), C_cooling(t)}',3)
    p('Equation (3) assumes that usable power and cooling subsystems can support the common workload even when their selected banks differ. It is not a rack-level dependency graph. The evaluated scope excludes the optional server, operating-system, network, storage, and application abstractions available in the software, keeping the experiment aligned with infrastructure redundancy alone.')
    fig('fig2_topology','Fig. 2. Capacity abstraction for the three architectures. Selection of a complete 2N train is not equivalent to pooling all installed units.')

    h('B. Failure and Repair Processes',True)
    p('Every active component starts healthy at t=0. Independent operating lifetimes are exponentially distributed with annual rate λ_g and hourly mean 8760/λ_g. Repair duration r_g is fixed at the selected MTTR. For successive failures of a component, the renewal process is')
    eq('X_g ∼ Exp(λ_g / 8760)',4)
    eq('t_(k+1) = t_k + r_g + X_(g,k+1)',5)
    p('A component therefore does not acquire a second independent operating failure while it is already under repair. Perfect repair returns it to the healthy state. Different components may fail during one another’s repair windows, producing natural overlaps. Failure starts and repair completions are sorted chronologically, with repairs processed before starts at an identical timestamp. Events at the trial horizon contribute no duration.')
    p('The event sweep retains a count of active failure causes for each unit. A repair clears its corresponding cause, but does not restore a unit that is still unavailable due to another cause, such as a shared hazard. This rule prevents a background repair from prematurely ending a common-cause outage. The independent renewal streams are not suspended by externally imposed shared hazards; cause overlap is a simplifying assumption of the scenario model.')

    h('C. Performance Measures',True)
    p('For trial i with horizon T hours, downtime D_i counts any interval in which the supported workload is below L. Unserved energy E_i measures the severity as well as the duration of the deficit:')
    eq('D_i = ∫₀ᵀ 1{C_i(t) < L} dt',6)
    eq('E_i = ∫₀ᵀ [L − C_i(t)] dt',7)
    eq('A_i = 1 − D_i / T',8)
    p('The exact piecewise-constant integration avoids adding overlapping outages twice. For a Y-year trial, annualized downtime is 60D_i/Y minutes and annualized unserved energy is E_i/Y kWh. Here Y=1 and T=8760. The SLA breach indicator is 1{A_i<0.99982}; its one-year downtime budget is 94.608 minutes. Mean availability, probability of any interruption, and probability of exceeding this budget are distinct outputs and are not interchangeable.')
    fig('fig3_overlap','Fig. 3. Deterministic N+1 example: failures on [1,5) and [3,7) h cause a 2 h service deficit, not 8 h of downtime. Served load is capped at 10 MW; the lost energy is 5 MWh.')

    h('D. Shared Hazards and Paired Trials',True)
    p('A common-cause scenario adds Poisson opportunities with rate ν per year, each accepted independently with probability p. The accepted hazard rate is νp. A accepted event disables every UPS and generator in bank A or across both banks for a fixed interval r_cc. The present experiments use ν=1/year, p=0.10, and r_cc=4 h. These are declared scenario assumptions, not observed site statistics. The added process changes marginal failure risk; it does not hold marginal probabilities fixed while varying a correlation coefficient.')
    p('Random-number streams are indexed by seed, component type, bank, unit, batch offset, and renewal ordinal. Corresponding units in different architectures share draws, while added units receive distinct streams. Reusing streams in a rate-sensitivity experiment reduces avoidable Monte Carlo noise in differences. Different master seeds generate separate replicates. The event generator uses fixed batches of 256 trials so extending a run preserves its existing prefix and bounds event-list memory.')

    h('IV. Implementation and Verification')
    p('The application combines a React/TypeScript frontend with a Python FastAPI backend and NumPy/SciPy simulation routines. Configuration validation precedes topology expansion. Simulation and comparison endpoints return trial-level records, aggregate statistics, input settings, and version identifiers; manual injection and replay use the same capacity evaluator as the simulation. Exported data are therefore inspectable rather than only rendered as dashboard scores. No sensors, live telemetry, machine-learning model, or production DCIM integration is required.')
    p('For each trial, the engine initializes healthy counts, generates operating and optional shared events, sweeps event times, integrates (6) and (7), and records the final metrics. For m events, sorting has O(m log m) time complexity; capacity updates then proceed incrementally. No time-step discretization is used. The web implementation also limits job size and concurrent CPU work, but the paper executes the engine directly to avoid conflating browser performance with numerical behaviour.')
    table('I','Deterministic verification cases',['Case','Expected result'],[
        ['Healthy, all architectures','10,000 kW served'],
        ['N: one UPS unavailable','7,500 kW served'],
        ['N+1: one UPS unavailable','10,000 kW served'],
        ['N+1: two UPS unavailable','7,500 kW served'],
        ['2N: UPS-A01 + PDU-B01','7,500 kW served'],
        ['N: outages [1,5), [3,7)','6 h union downtime'],
        ['N+1: Fig. 3 overlap','2 h; 5,000 kWh lost'],
        ['Repair and failure at t=3','No positive-duration overlap']
    ],[1.85,1.60])
    p('All 41 existing regression tests passed in the manuscript preparation environment. Tests cover topology sizing, identical demand, disconnected trains, cause reference counts, event clipping, paired streams, reproducibility, API validation, and risk bounds. Table I gives hand-checkable cases. The dual-train split-path case is particularly important: pooling all healthy UPSs and PDUs would incorrectly classify it as fully available.')
    p('A separate regression check samples a single generator renewal process and compares its mean downtime with the small-unavailability approximation λr. Because λ is measured per year and r in hours, λr already has units of hours/year. Multiplying it by 8760 again would be dimensionally incorrect. The exact stationary component unavailability is')
    eq('q_g = (λ_g r_g) / (8760 + λ_g r_g)',9)
    p('The 100,000-realization renewal check agrees with λr within its predefined 2.5% tolerance. It verifies sampling and unit conversion, not the validity of the underlying generator rate. A healthy initial condition makes the one-year experiment transient; stationary calculations are used as sanity checks, not as identical finite-horizon ground truth.')

    h('V. Input Data and Experimental Design')
    h('A. Provenance and Scope of the Input Dataset',True)
    p('Table II contains every failure-rate class active in the study. Three rows are transcribed from Table 2 of Wiboonrat [4], which cites IEEE Std 493 [3]. They are secondary-source records, not direct verification of the original standard. The UPS entry describes a small computer-room unit; the generator entry describes continuous-duty equipment in the 250 kW–1.5 MW range. Applying these statistics to the larger modules used here is an explicit extrapolation.')
    table('II','Active failure-rate inputs',['Class','λ (/year)','MTTR (h)','Provenance'],[
        ['UPS','0.010','2.00','[4], E39-200'],
        ['Generator','0.580','25.74','[4], E18-121'],
        ['ATS','0.030','1.64','[4], E34-110'],
        ['PDU','0.020','3.00','Illustrative'],
        ['CRAC','0.055','5.50','Illustrative']
    ],[.72,.69,.72,1.32],note='PDU and CRAC values are scenario assumptions from the supplied project specification; neither is verified by the cited table. All repair times are fixed in this experiment.')
    p('The generator row has an internal discrepancy: its printed rate is 0.58/year, whereas 115 failures over 266 unit-years imply 0.432330827/year. Its printed MTBF of 15,033.8 h implies approximately 0.582687/year. The baseline retains the printed rate, and a sensitivity experiment substitutes the count/exposure ratio. Neither interpretation resolves the source inconsistency. Reporting both is preferable to silently treating one as validated evidence.')
    p('The PDU and cooling inputs remain illustrative. Consequently, all numerical findings below are conditional scenario results. They must not be advertised as measured IEEE availability, validated facility forecasts, or estimates applicable to a particular manufacturer. This distinction preserves the usefulness of an executable methods study without concealing the data still needed for engineering deployment.')

    h('B. Controlled Configuration',True)
    p('Demand is fixed at L=10,000 kW. Unit ratings and the resulting counts are listed in Table III. All three designs serve the same load; installed capacity varies as the intended consequence of redundancy. Generators are continuously required in islanded mode. This is not an estimate for a grid-connected site whose generators operate only during utility outages. Cooling capacity is represented as supported IT kW, with no thermal delay.')
    table('III','Fixed ratings and installed unit counts',['Group','Unit (kW)','N','N+1','2N'],[
        ['UPS','2,500','4','5','8'],['Generator','5,000','2','3','4'],
        ['PDU','2,500','4','5','8'],['ATS','10,000','1','2','2'],
        ['CRAC','2,000','5','6','10'],['Total units','—','16','21','32'],
        ['Count index','—','1.000','1.313','2.000']
    ],[.88,.83,.52,.64,.58])
    p('Each architecture–scenario combination uses three master seeds (42, 2026, and 745), with 10,000 independent one-year trials per seed. Thus each reported combination contains 30,000 facility-years. Five scenarios and three architectures produce 450,000 facility-years in total. Shared streams pair architectures within each seed; the three seeds provide a check on seed-specific variation. No maintenance, IT-service faults, or diagnostic reruns are mixed into these comparisons.')
    table('IV','Experimental scenarios',['Scenario','Change from nominal'],[
        ['Nominal','Independent operating hazards'],
        ['20× stress','All five λ values multiplied by 20'],
        ['Bank-A shared hazard','νp = 0.10/year; 4 h; A only'],
        ['Site-wide shared hazard','νp = 0.10/year; 4 h; all banks'],
        ['Generator source audit','λ_GEN = 115/266 per year']
    ],[1.26,2.19],note='Shared hazards affect UPS and generator units only. Their rate is not multiplied by the independent-hazard stress factor; stress and shared-hazard scenarios are separate experiments.')
    p('The 20× experiment accelerates operating hazards to generate more overlapping outages and probe topology behaviour. Repair duration is unchanged. It is a deliberately adverse synthetic condition, not a twenty-year forecast or a weather model. The bank and site scenarios use identical accepted-event streams, isolating the scope of affected units rather than changes in incident timing.')

    h('C. Statistical Reporting and Reproducibility',True)
    p('Reported means pool the 30,000 trials of each combination. Approximate two-sided 95% mean intervals use the sample standard deviation and 1.96 standard errors. Comparisons use per-trial paired downtime differences, rather than adding independent error bars. Binomial SLA and any-outage probabilities use two-sided 95% Clopper–Pearson intervals [8]. The mean intervals quantify simulation sampling error only; they omit uncertainty in failure rates, repair distributions, and model structure.')
    eq('CI_mean ≈ D̄ ± 1.96 s_D / √M',10)
    p('Rare-event evidence must be interpreted separately from the nominal trial count. If no outage appears in M trials, an exact one-sided 95% upper bound for the probability of any outage is 1−0.05^(1/M), approximately 3/M [9]. This is not a direct estimate of mean downtime. The software uses a conservative risk-derived bound for zero-outage samples; the study does not claim perfect reliability from an empty outage sample.')
    p('Runs use engine 2.3.0 and dataset 2026-09-12-audited-assumptions-v2. The supplementary archive contains the experiment script, input configurations, version and hash metadata, compressed per-run JSON, pooled CSV, and figure-generation code. Recording a seed without recording the engine, data, and coupling convention would not be sufficient for exact reproduction.')

    h('VI. Results and Analysis')
    h('A. Independent-Hazard Baseline',True)
    table('V','Nominal independent-hazard results',['Metric','N','N+1','2N'],[
        ['Availability (%)']+[val('baseline',a,'availability_percent',5) for a in A],
        ['Downtime (min/y)']+[val('baseline',a) for a in A],
        ['95% mean CI']+[f"{R[('baseline',a)]['ci_low']:.2f}–{R[('baseline',a)]['ci_high']:.2f}" for a in A],
        ['Outage trials / 30k']+[str(R[('baseline',a)]['outage_trials']) for a in A],
        ['SLA breach (%)']+[f"{100*R[('baseline',a)]['sla_probability']:.3f}" for a in A],
        ['Unserved (kWh/y)']+[val('baseline',a,'unserved_kwh',1) for a in A]
    ],[1.16,.78,.76,.75],note='Mean CIs are approximate. The SLA benchmark is 99.982% over a one-year trial, not a Tier certification criterion. All values are model-dependent simulation outputs.')
    p('Table V shows the expected large improvement from adding redundancy to an unprotected capacity baseline. Mean annual downtime falls from '+val('baseline','N')+' min for N to '+val('baseline','N+1')+' min for N+1 and '+val('baseline','2N')+' min for 2N. Relative to N, these reductions are '+f"{100*(1-R[('baseline','N+1')]['downtime_minutes']/R[('baseline','N')]['downtime_minutes']):.2f}"+'% and '+f"{100*(1-R[('baseline','2N')]['downtime_minutes']/R[('baseline','N')]['downtime_minutes']):.2f}"+'%, respectively, under the nominal assumptions.')
    paired=next(z for z in S['paired'] if z['scenario']=='baseline' and z['left']=='N+1')
    p('The comparison between redundant designs is more subtle. The paired difference D_N+1−D_2N is '+f"{paired['reduction']:.2f}"+' min/year, with an approximate 95% interval ['+f"{paired['low']:.2f}, {paired['high']:.2f}"+']. Because the interval includes zero, the nominal study does not resolve a difference in mean downtime between N+1 and 2N. Ranking them from their point estimates alone would overstate the evidence.')
    p('The fraction of trials breaching the SLA is not implied by the mean availability. The exact 95% breach-probability intervals are '+', '.join(a+': '+f"[{100*R[('baseline',a)]['sla_ci'][0]:.3f}, {100*R[('baseline',a)]['sla_ci'][1]:.3f}]%" for a in A)+'. A design may have a mean below its downtime budget and still experience rare years with substantial interruption. Unserved energy additionally distinguishes a partial capacity deficit from a full-site outage of the same duration.')
    fig('fig4_comparison','Fig. 4. Nominal and 20× independent-hazard scenarios. Error bars are approximate 95% mean intervals; the logarithmic axis reveals differences hidden by availability percentages near 100%.')

    h('B. Why More Installed Capacity Need Not Dominate',True)
    p('At 20× operating hazards, mean annual downtime is '+val('stress20','N')+', '+val('stress20','N+1')+', and '+val('stress20','2N')+' min for N, N+1, and 2N. The higher outage count makes the distinction between pooled spares and disconnected trains easier to observe. This does not contradict redundancy engineering: the compared architectures have different service-success rules, not simply different numbers of interchangeable units.')
    stresspair=next(z for z in S['paired'] if z['scenario']=='stress20' and z['left']=='N+1')
    p('Under stress, D_N+1−D_2N is '+f"{stresspair['reduction']:.2f}"+' min/year, with approximate 95% interval ['+f"{stresspair['low']:.2f}, {stresspair['high']:.2f}"+']. For these assumptions, N+1 therefore has lower mean downtime than disconnected 2N. This finding is conditional on ideal spare sharing within a group, complete-train selection without cross-ties, and continuously exposed redundant equipment.')
    p('A simple generator-only calculation explains the direction. With two required generators and per-generator unavailability q, N+1 has three generators and loses service only when at least two fail, giving 3q²−2q³. A 2N design has two disconnected two-generator trains; each train needs both generators, so both trains are impaired with probability (2q−q²)². At small q, the leading terms are 3q² and 4q², respectively. Other power-group failures can also disable opposite trains. This reasoning is an architectural check, not a replacement for the complete simulation.')

    h('C. Bank-Local Versus Site-Wide Dependence',True)
    table('VI','Shared-hazard and source-audit downtime',['Scenario','N','N+1','2N'],[
        ['Bank-A hazard']+[val('shared_bank',a) for a in A],
        ['Site-wide hazard']+[val('shared_site',a) for a in A],
        ['Alternative GEN λ']+[val('generator_alternative',a) for a in A]
    ],[1.26,.73,.73,.73],note='Units: mean minutes/year; 30,000 trials per cell. Hazards and alternative generator rate are defined in Tables II and IV. These are separate perturbations, not combined scenarios.')
    p('A bank-A hazard largely defeats spare capacity confined to the same bank. N+1 downtime rises to '+val('shared_bank','N+1')+' min/year, whereas 2N remains at '+val('shared_bank','2N')+' min/year because its intact B train usually supports the load. This is the failure-domain separation that the independent-hazard experiment cannot capture by counting spare units alone.')
    p('When the same accepted incident affects both banks, 2N downtime rises to '+val('shared_site','2N')+' min/year. The site-wide N+1 result is unchanged from its bank-local result because N+1 has only one bank. At an accepted rate of 0.10/year and 4 h recovery, the added full-site hazard has a first-order contribution of 24 min/year before accounting for overlaps and boundary clipping. Fig. 5 demonstrates why redundant capacity cannot compensate for a hazard that removes every usable power train.')
    fig('fig5_dependence','Fig. 5. Shared power-hazard scope changes the value of dual trains. These scenarios add an assumed hazard; they are not fixed-marginal correlation experiments.')
    p('An architectural recommendation must therefore name the hazard domain. Physical separation helps against a bank-local event, but not necessarily a shared site fault. Conversely, a pooled spare can be efficient for independent unit failures. The results support neither a universal N+1 preference nor a universal 2N preference; they support evaluation against explicit, plausible failure mechanisms.')

    h('D. Source Sensitivity and Seed Variation',True)
    p('Replacing the printed generator rate of 0.58/year with 115/266 per year changes mean downtime to '+val('generator_alternative','N')+', '+val('generator_alternative','N+1')+', and '+val('generator_alternative','2N')+' min/year. Figure 6 shows that a single unresolved source entry materially changes the absolute estimates. Because redundancy relies on overlapping events, its response is not generally linear in an individual component rate.')
    fig('fig6_source_audit','Fig. 6. Generator-source sensitivity with all other inputs fixed. Error bars are approximate 95% mean intervals. Neither rate interpretation is independently validated for the configured equipment.')
    p('The nominal per-seed mean downtime ranges are '+f"{min(R[('baseline','N+1')]['seed_means']):.2f}–{max(R[('baseline','N+1')]['seed_means']):.2f}"+' min/year for N+1 and '+f"{min(R[('baseline','2N')]['seed_means']):.2f}–{max(R[('baseline','2N')]['seed_means']):.2f}"+' min/year for 2N. Figure 7 plots cumulative estimates from each seed. Their visible variation reinforces why a single seeded result with many printed decimal places is not an adequate reliability claim. Three seeds provide a diagnostic, not a guarantee that all rare-event behaviour has been resolved.')
    fig('fig7_convergence','Fig. 7. Cumulative nominal downtime estimates for three independent seeds. The redundant cases remain sensitive to relatively few interruption years even as trial count increases.')

    h('E. Infrastructure Overhead',True)
    p('The equal-weight installed-unit index is 1.000 for N, 1.3125 for N+1, and 2.000 for 2N. Thus the model adds 31.25% and 100% more units, respectively. This is an inventory proxy, not a cost estimate: one cooling unit, generator, and ATS do not have equal prices or capacities. The application’s editable budget inputs are not verified quotations, so monetary totals and return-on-investment claims are deliberately excluded from the paper.')
    p('Under the nominal model, N+1 attains a large reduction relative to N with fewer installed units than 2N. Under a bank-local shared hazard, the separate train provides an advantage that the count index alone misses. Procurement decisions therefore require a weighted equipment bill, separation and switching costs, installation, staffing, fuel, maintenance logistics, energy use, and the economic consequences of interruption. The present study quantifies a conditional availability–inventory comparison rather than claiming an optimal investment.')

    h('VII. Discussion, Limitations, and Future Work')
    p('The primary limitation is external validity. The study is not calibrated against facility incident logs, and two active classes use illustrative rates. Secondary-source equipment classes do not match every configured capacity, the generator record is inconsistent, and rounded source rates discard statistical detail. Consequently, tight sampling intervals cannot justify equally tight confidence in real-facility availability. Future work must replace illustrative inputs with equipment-specific evidence and propagate input uncertainty through the comparison.')
    p('The model also idealizes repair and operation. Constant MTTR excludes heavy-tailed repair times, staff and spare-part constraints, and multi-unit repair queues. Exponential operating lifetimes exclude ageing, infant mortality, and load-dependent hazards. All units start healthy, so long-run estimates should eventually be checked with equilibrium initialization or a warm-up study. Ideal failover omits switching interruption, and the capacity model does not represent protection coordination, physical rack mapping, thermal inertia, or cooling recovery transients.')
    p('Islanded operation is a deliberate scope choice. A grid-connected model requires an explicit utility-outage process, battery autonomy, generator start failure per demand, retry behaviour, fuel exhaustion, and transfer logic. A continuous-duty failure rate cannot substitute for a start-failure probability. Similarly, naturally overlapping independent failures are not a general cascade model. A realistic cascade needs mechanisms that alter downstream state or hazard after an initiating fault.')
    p('Shared hazards here are simple additive stressors. Their frequency and duration are not inferred from observed incidents, and shared failures do not suspend underlying operating clocks. These assumptions should be tested against dependency models with calibrated marginals and recovery dynamics. Scheduled maintenance can be injected by the implementation, but it is excluded from the reported stochastic study to keep causal comparisons interpretable. Neither this omission nor deterministic maintenance checks establish concurrent maintainability or fault-tolerance certification.')
    p('Statistical limitations include skewed downtime distributions, approximate normal intervals for means, and a finite set of seeds and scenarios. More trials reduce sampling error but not model bias. Rare-event techniques such as importance sampling may improve efficiency if their estimators are independently validated. Future studies should also vary load margin, unit granularity, repair distributions, and cross-tie rules rather than assuming that conclusions from one 10 MW configuration generalize to every facility.')

    h('VIII. Conclusion')
    p('DC-Resilience provides an inspectable Monte Carlo workflow for comparing data-centre redundancy under explicit capacity and failure assumptions. Across 450,000 simulated facility-years, both redundant designs substantially reduce nominal downtime relative to N, while their nominal difference remains statistically unresolved. The accelerated-hazard scenario favours pooled N+1 capacity under the model’s disconnected-train rules; bank-local shared failures favour separate 2N trains; site-wide hazards compromise both. An unresolved generator-source discrepancy materially shifts absolute results.')
    p('The central result is methodological: redundancy labels and installed-unit counts are insufficient without topology, hazard scope, data provenance, and uncertainty. The framework supplies a reproducible basis for such comparisons, but the present scenario study is not a validated facility forecast or Tier assessment. Equipment-specific data verification, empirical dependency calibration, and richer operational models are required before using its numerical outputs for facility design commitments.')

    h('Data and Code Availability')
    p('The supplementary material includes 45 compressed run records, pooled summaries, original SVG/PNG figures, and the scripts paper/run_study.py and paper/build_paper.py. Source: github.com/RomitDeokar/DC-Resilience. The study manifest records engine and dataset versions, source hashes, seeds, and execution metadata. All reported results can be recomputed from the supplied run records; no proprietary incident log or field measurement is claimed.',size=9)

    h('References')
    refs=[
        '[1] Uptime Institute, “Tier Classification System.” [Online]. Available: https://uptimeinstitute.com/tiers. Accessed: Sep. 12, 2026.',
        '[2] K. Heslin, “Myths and misconceptions regarding the Uptime Institute’s Tier certification system,” Uptime Institute Journal, Jul. 20, 2017. [Online]. Available: https://journal.uptimeinstitute.com/myths-and-misconceptions-regarding-the-uptime-institutes-tier-certification-system/.',
        '[3] IEEE Recommended Practice for the Design of Reliable Industrial and Commercial Power Systems, IEEE Std 493-2007, 2007.',
        '[4] M. Wiboonrat, “Condition-based maintenance for data center operations management,” in Operations Management—Emerging Trend in the Digital Era, IntechOpen, online Oct. 26, 2020; book ed., 2021, doi: 10.5772/intechopen.93945.',
        '[5] Y. Lei and A. Q. Huang, “Data center power distribution system reliability analysis tool based on Monte Carlo next event simulation method,” in Proc. IEEE Energy Conversion Congress and Exposition (ECCE), 2017, pp. 2031–2035, doi: 10.1109/ECCE.2017.8096406.',
        '[6] D. Ford, F. Labelle, F. I. Popovici, M. Stokely, V.-A. Truong, L. Barroso, C. Grimes, and S. Quinlan, “Availability in globally distributed storage systems,” in Proc. 9th USENIX Symp. Operating Systems Design and Implementation (OSDI), 2010.',
        '[7] T. M. Julitz, A. Tordeux, N. Schlüter, and M. Löwer, “Reliability of redundant M-out-of-N architectures with dependent components: A comprehensible approach with Monte Carlo simulation,” arXiv:2402.18187, 2024, doi: 10.48550/arXiv.2402.18187.',
        '[8] C. J. Clopper and E. S. Pearson, “The use of confidence or fiducial limits illustrated in the case of the binomial,” Biometrika, vol. 26, no. 4, pp. 404–413, 1934, doi: 10.1093/biomet/26.4.404.',
        '[9] J. A. Hanley and A. Lippman-Hand, “If nothing goes wrong, is everything all right? Interpreting zero numerators,” JAMA, vol. 249, no. 13, pp. 1743–1745, 1983, doi: 10.1001/jama.1983.03330370053031.'
    ]
    for ref in refs:
        z=p(ref,size=8,indent=False)
        z.paragraph_format.left_indent=Inches(.18);z.paragraph_format.first_line_indent=Inches(-.18)
        z.paragraph_format.space_after=Pt(3)
    # A continuous final section balances the last pair of text columns.
    last=D.add_section(WD_SECTION_START.CONTINUOUS)
    cols=last._sectPr.find(qn('w:cols'));cols.set(qn('w:num'),'1')
    D.paragraphs[-1].paragraph_format.space_after=Pt(0)
    D.paragraphs[-1].paragraph_format.line_spacing=Pt(1)
    D.save(ROOT/'DC_Resilience_IEEE_Conference_Paper.docx')
    text='\n'.join(z.text for z in D.paragraphs)
    (ROOT/'results/manuscript_text.txt').write_text(text)
    print('Saved manuscript; body words:',len(text.split()),'tables:',len(D.tables),'figures:',len(D.inline_shapes))

if __name__=='__main__':
    figures()
    manuscript()
